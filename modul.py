from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_lab_manual():
    # Inisialisasi Dokumen
    doc = Document()

    # ===========================
    # HALAMAN JUDUL
    # ===========================
    title = doc.add_heading('MODUL PRAKTIKUM VIRTUAL LABORATORY', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Panduan Lengkap Simulasi Fisika & Analisis Data')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()

    # ===========================
    # DAFTAR ISI (Placeholder)
    # ===========================
    doc.add_heading('DAFTAR ISI', level=1)
    doc.add_paragraph('1. Pendahuluan')
    doc.add_paragraph('2. Simulasi Gerak Parabola (Pygame)')
    doc.add_paragraph('3. Simulasi Revolusi Bumi (Matplotlib)')
    doc.add_paragraph('4. Analisis Data Statistik (SEM/PLS)')
    doc.add_paragraph('5. AI Grader untuk Evaluasi Otomatis')
    doc.add_page_break()

    # ===========================
    # BAB 1: PENDAHULUAN
    # ===========================
    doc.add_heading('BSB 1: PENDAHULUAN', level=1)
    doc.add_paragraph(
        'Modul ini dirancang untuk memandu mahasiswa dan peneliti dalam menggunakan '
        'alat-alat virtual laboratory yang telah dikembangkan. Setiap modul mencakup '
        'tujuan pembelajaran, prasyarat sistem, dan langkah-langkah penggunaan.'
    )

    # ===========================
    # BAB 2: GERAK PARABOLA
    # ===========================
    doc.add_heading('BAB 2: SIMULASI GERAK PARABOLA', level=1)
    
    doc.add_heading('A. Deskripsi', level=2)
    doc.add_paragraph(
        'Simulasi ini memvisualisasikan gerak peluru secara real-time menggunakan '
        'engine Pygame. Pengguna dapat mengubah sudut dan kecepatan awal.'
    )

    doc.add_heading('B. Cara Menjalankan', level=2)
    p = doc.add_paragraph()
    p.add_run('1. Buka terminal atau command prompt.\n')
    p.add_run('2. Jalankan perintah: python parabola_pygame.py\n').bold = True
    p.add_run('3. Gunakan keyboard untuk interaksi.')

    doc.add_heading('C. Tugas Praktikum', level=2)
    doc.add_paragraph(
        '1. Variasikan sudut tembak dari 30 hingga 60 derajat.\n'
        '2. Catat jarak maksimum yang dicapai peluru.\n'
        '3. Bandingkan hasil simulasi dengan perhitungan teoritis.'
    )

    # ===========================
    # BAB 3: REVOLUSI BUMI
    # ===========================
    doc.add_heading('BAB 3: SIMULASI REVOLUSI BUMI', level=1)
    
    doc.add_heading('A. Deskripsi', level=2)
    doc.add_paragraph(
        'Visualisasi orbit bumi mengelilingi matahari menggunakan Matplotlib Animation (3D). '
        'Menunjukkan konsep kemiringan sumbu bumi dan pergantian musim.'
    )

    doc.add_heading('B. Cara Menjalankan', level=2)
    p = doc.add_paragraph()
    p.add_run('1. Jalankan perintah: python revolusi_bumi.py\n').bold = True
    p.add_run('2. Jendela plot akan muncul secara otomatis.\n')
    p.add_run('3. Perhatikan label bulan dan posisi bumi.')

    # ===========================
    # BAB 4: ANALISIS SEM/PLS
    # ===========================
    doc.add_heading('BAB 4: ANALISIS STATISTIK (SEM/PLS)', level=1)
    
    doc.add_heading('A. Persiapan Data', level=2)
    doc.add_paragraph(
        'Pastikan data Anda dalam format CSV dengan pemisah titik koma (;) '
        'jika menggunakan format angka Eropa (koma sebagai desimal).'
    )

    doc.add_heading('B. Langkah Analisis', level=2)
    doc.add_paragraph(
        '1. Validasi Data: Periksa Normalitas dan Multikolinearitas (VIF).\n'
        '2. Model Struktural: Definisikan hubungan antar variabel laten.\n'
        '3. Evaluasi Model: Periksa R-Squared dan Path Coefficients.'
    )

    # ===========================
    # BAB 5: AI GRADER
    # ===========================
    doc.add_heading('BAB 5: AI EDUCATIONAL TOOL', level=1)
    doc.add_heading('A. Konfigurasi API', level=2)
    doc.add_paragraph(
        'Alat ini menggunakan Google Generative AI (Gemini). Anda harus memasukkan API Key '
        'pada variabel "API_KEY" di dalam skrip.'
    )
    doc.add_paragraph(
        'PERINGATAN: Jangan membagikan API Key Anda secara publik.', style='Intense Quote'
    )

    # ===========================
    # BAB 6: TAMPILAN AWAL SIMULASI
    # ===========================
    doc.add_heading('BAB 6: TAMPILAN AWAL SIMULASI', level=1)
    
    doc.add_heading('A. Deskripsi', level=2)
    doc.add_paragraph(
        'Pada bagian ini, Anda akan diperkenalkan dengan antarmuka awal dari aplikasi Virtual Lab. '
        'Tampilan ini akan muncul setelah Anda menjalankan skrip simulasi.'
    )

    doc.add_heading('B. Elemen-Elemen Kunci', level=2)
    doc.add_paragraph(
        '1. Menu Utama: Akses cepat ke berbagai modul praktikum.\n'
        '2. Panel Simulasi: Area untuk menjalankan dan mengamati simulasi.\n'
        '3. Laporan Hasil: Menampilkan hasil perhitungan dan analisis.'
    )

    doc.add_heading('C. Screenshot', level=2)
    if os.path.exists('screenshots/nama_file_screenshot.png'):
        doc.add_picture('screenshots/nama_file_screenshot.png', width=Inches(4.5))
    else:
        doc.add_paragraph('Screenshot belum tersedia.')

    # ===========================
    # BAB X: SIMULASI ...
    # ===========================
    doc.add_heading('BAB X: SIMULASI ...', level=1)

    doc.add_heading('A. Deskripsi', level=2)
    doc.add_paragraph('Simulasi ini ...')

    doc.add_heading('B. Tujuan', level=2)
    doc.add_paragraph('1. ...')

    doc.add_heading('C. Prasyarat', level=2)
    doc.add_paragraph('1. Python 3.x ...')

    doc.add_heading('D. Langkah Menjalankan', level=2)
    doc.add_paragraph('1. Buka terminal ...')

    doc.add_heading('E. Tugas Praktikum', level=2)
    doc.add_paragraph('1. ...')

    doc.add_heading('F. Tampilan Awal Simulasi', level=2)
    if os.path.exists('Virtual Lab/screenshots/nama_file_screenshot.png'):
        doc.add_picture('Virtual Lab/screenshots/nama_file_screenshot.png', width=Inches(4.5))
        doc.add_paragraph('Gambar: Tampilan awal aplikasi ...')
    else:
        doc.add_paragraph('Screenshot belum tersedia.')

    doc.add_heading('G. Catatan/Analisis', level=2)
    doc.add_paragraph('Analisis hasil simulasi ...')

    # Simpan file
    file_name = 'Modul_Praktikum_Virtual_Lab.docx'
    doc.save(file_name)
    print(f"Modul berhasil dibuat: {file_name}")

if __name__ == "__main__":
    create_lab_manual()